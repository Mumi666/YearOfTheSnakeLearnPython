import datetime
import openpyxl
from docx import Document
from pathlib import Path

today = datetime.date.today().strftime('%Y-%m-%d')

excel_file = './file/customer.xlsx'
template_doc = './file/邀请函模版.docx'

format_content = {
    '<姓名>': 'no_name',
    '<性别>': 'no_gender',
    '<今天日期>': 'no_date',
}

def generate_invite_doc():
    doc = Document(template_doc)
    # 替换模版中的关键字
    for paragraph in doc.paragraphs:
        for key, value in format_content.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, value)
    result_dir = Path('./result/')
    result_dir.mkdir(parents=True, exist_ok=True)

    save_file_path = result_dir / f"{format_content['<姓名>']}.docx"
    doc.save(save_file_path)


def load_excel_customer(excel_file_path):
    # 读取excel中的内容
    wb = openpyxl.load_workbook(excel_file_path)
    sheet = wb.active

    # 遍历工作表中的每一行
    for row in sheet.iter_rows(min_row=2):  # 从第2行开始，跳过表头
        # 修改字典中的值
        format_content['<姓名>'] = row[0].value  # A列是姓名
        format_content['<性别>'] = row[1].value  # B列是性别
        format_content['<今天日期>'] = today

        # 调用generate_invite_doc()
        generate_invite_doc()

if __name__ == '__main__':
    load_excel_customer(excel_file)
